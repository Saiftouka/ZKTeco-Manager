from datetime import datetime
from io import BytesIO
from time import perf_counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app.services.correction_pointage_service import appliquer_corrections_pointages
from app.services.pointage_cache_service import assurer_pointages_intervalle


def _dt(valeur):
    if isinstance(valeur, datetime):
        return valeur

    texte = str(valeur or "").strip()
    if not texte:
        raise ValueError("La date et l'heure sont obligatoires.")

    formats = (
        "%Y-%m-%dT%H:%M",
        "%Y-%m-%dT%H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y-%m-%d %H:%M:%S",
        "%d/%m/%Y %H:%M",
        "%d/%m/%Y %H:%M:%S",
    )

    for format_datetime in formats:
        try:
            return datetime.strptime(texte, format_datetime)
        except ValueError:
            continue

    try:
        return datetime.fromisoformat(texte)
    except ValueError as exc:
        raise ValueError("Format de date/heure invalide.") from exc


def _nom_fichier(valeur):
    texte = str(valeur or "pointeuse").strip()
    for caractere in '\\/:*?"<>|':
        texte = texte.replace(caractere, "_")
    return texte.replace(" ", "_") or "pointeuse"


def _entier(valeur):
    if valeur is None or valeur == "":
        return None
    try:
        return int(valeur)
    except (TypeError, ValueError):
        return None


def _normaliser(pointage):
    if not isinstance(pointage, dict):
        return None

    timestamp_source = pointage.get("timestamp")

    if timestamp_source in (None, ""):
        date_texte = str(pointage.get("date", "") or "").strip()
        heure_texte = str(pointage.get("heure", "") or "").strip()
        timestamp_source = (date_texte + " " + heure_texte).strip()

    try:
        timestamp = _dt(timestamp_source)
    except Exception:
        return None

    return {
        "uid": _entier(pointage.get("uid")),
        "matricule": str(pointage.get("matricule", "") or "").strip(),
        "nom": str(pointage.get("nom", "") or "").strip(),
        "badge": str(pointage.get("badge", "") or "").strip(),
        "timestamp": timestamp,
        "date": timestamp.strftime("%d/%m/%Y"),
        "heure": timestamp.strftime("%H:%M:%S"),
        "status": _entier(pointage.get("status")),
        "punch": _entier(pointage.get("punch")),
    }


def _cle_tri(pointage):
    matricule = str(pointage.get("matricule", "") or "").strip()
    timestamp = pointage.get("timestamp")
    if not isinstance(timestamp, datetime):
        timestamp = datetime.max

    if matricule.isdigit():
        cle_matricule = (0, int(matricule), "")
    else:
        cle_matricule = (1, 0, matricule.casefold())

    return cle_matricule + (timestamp,)


def creer_excel_pointages(
    pointeuse_id,
    nom_pointeuse,
    ip,
    port,
    date_heure_debut,
    date_heure_fin,
):
    debut_total = perf_counter()

    try:
        debut = _dt(date_heure_debut)
        fin = _dt(date_heure_fin)
    except Exception as erreur:
        return {
            "success": False,
            "contenu": None,
            "nom_fichier": None,
            "nombre": 0,
            "nombre_terminal": 0,
            "erreur": str(erreur),
        }

    if debut > fin:
        return {
            "success": False,
            "contenu": None,
            "nom_fichier": None,
            "nombre": 0,
            "nombre_terminal": 0,
            "erreur": "La date de debut doit etre anterieure ou egale a la date de fin.",
        }

    # ========================================================
    # SQLITE-FIRST
    # ========================================================

    mesure = perf_counter()

    try:
        local = assurer_pointages_intervalle(
            pointeuse_id=pointeuse_id,
            ip=ip,
            port=port,
            debut=debut,
            fin=fin,
        )
    except Exception as erreur:
        return {
            "success": False,
            "contenu": None,
            "nom_fichier": None,
            "nombre": 0,
            "nombre_terminal": 0,
            "erreur": "Impossible de lire l'historique local : " + str(erreur),
        }

    temps_sqlite = perf_counter() - mesure

    if not isinstance(local, dict):
        return {
            "success": False,
            "contenu": None,
            "nom_fichier": None,
            "nombre": 0,
            "nombre_terminal": 0,
            "erreur": "Le service d'historique local a retourne une reponse invalide.",
        }

    if local.get("success") is not True:
        return {
            "success": False,
            "contenu": None,
            "nom_fichier": None,
            "nombre": 0,
            "nombre_terminal": 0,
            "erreur": local.get("erreur") or "Lecture SQLite impossible.",
        }

    pointages_bruts = local.get("pointages", [])
    if not isinstance(pointages_bruts, list):
        pointages_bruts = []

    pointages = []
    for pointage in pointages_bruts:
        normalise = _normaliser(pointage)
        if normalise is not None:
            pointages.append(normalise)

    # ========================================================
    # CORRECTIONS
    # ========================================================

    mesure = perf_counter()

    try:
        corrections = appliquer_corrections_pointages(
            pointeuse_id,
            pointages,
        )
    except Exception as erreur:
        return {
            "success": False,
            "contenu": None,
            "nom_fichier": None,
            "nombre": 0,
            "nombre_terminal": 0,
            "erreur": "Impossible d'appliquer les corrections : " + str(erreur),
        }

    temps_corrections = perf_counter() - mesure

    if corrections is None:
        corrections = {
            "success": True,
            "pointages": pointages,
            "nombre_corrections": 0,
            "erreur": None,
        }

    if not isinstance(corrections, dict):
        return {
            "success": False,
            "contenu": None,
            "nom_fichier": None,
            "nombre": 0,
            "nombre_terminal": 0,
            "erreur": "Le service de corrections a retourne une reponse invalide.",
        }

    if corrections.get("success") is not True:
        return {
            "success": False,
            "contenu": None,
            "nom_fichier": None,
            "nombre": 0,
            "nombre_terminal": 0,
            "erreur": corrections.get("erreur") or "Application des corrections impossible.",
        }

    pointages_corriges = corrections.get("pointages", pointages)
    if not isinstance(pointages_corriges, list):
        pointages_corriges = pointages

    pointages = []
    for pointage in pointages_corriges:
        if not isinstance(pointage, dict):
            continue

        ligne = dict(pointage)
        timestamp = ligne.get("timestamp")

        if not isinstance(timestamp, datetime):
            try:
                timestamp = _dt(timestamp)
            except Exception:
                continue

        ligne["timestamp"] = timestamp
        ligne["date"] = timestamp.strftime("%d/%m/%Y")
        ligne["heure"] = timestamp.strftime("%H:%M:%S")
        pointages.append(ligne)

    pointages.sort(key=_cle_tri)

    # ========================================================
    # INFORMATIONS DE SYNCHRONISATION
    # ========================================================

    synchronisation = local.get("synchronisation")
    if not isinstance(synchronisation, dict):
        synchronisation = {}

    try:
        nombre_terminal = int(synchronisation.get("nombre_terminal", 0) or 0)
    except (TypeError, ValueError):
        nombre_terminal = 0

    # ========================================================
    # EXCEL
    # ========================================================

    mesure = perf_counter()
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Pointages"
    feuille.sheet_view.showGridLines = False

    bleu = "17324D"
    orange = "F28C28"
    gris = "F4F7FA"
    ligne = "D9E2EA"
    blanc = "FFFFFF"
    texte = "172B3D"
    alternee = "F8FAFC"

    feuille.merge_cells("A1:D1")
    feuille["A1"] = "ZKTECO MANAGER | RAPPORT DE POINTAGES"
    feuille["A1"].font = Font(
        name="Aptos Display",
        size=18,
        bold=True,
        color=blanc,
    )
    feuille["A1"].fill = PatternFill(
        fill_type="solid",
        fgColor=bleu,
    )
    feuille["A1"].alignment = Alignment(
        horizontal="left",
        vertical="center",
    )
    feuille.row_dimensions[1].height = 34

    feuille.merge_cells("A2:D2")
    feuille["A2"] = "Rapport de pointages"
    feuille["A2"].font = Font(
        name="Aptos",
        size=10,
        italic=True,
        color="6B7A89",
    )

    informations = [
        ("Pointeuse", str(nom_pointeuse or "")),
        ("Adresse", f"{str(ip).strip()} : {int(port)}"),
        ("Du", debut),
        ("Au", fin),
        ("Pointages exportes", len(pointages)),
    ]

    for numero_ligne, (libelle, valeur) in enumerate(
        informations,
        start=4,
    ):
        cellule_libelle = feuille.cell(
            row=numero_ligne,
            column=1,
            value=libelle,
        )
        cellule_valeur = feuille.cell(
            row=numero_ligne,
            column=2,
            value=valeur,
        )

        cellule_libelle.font = Font(
            name="Aptos",
            bold=True,
            color=bleu,
        )
        cellule_libelle.fill = PatternFill(
            fill_type="solid",
            fgColor=gris,
        )
        cellule_valeur.font = Font(
            name="Aptos",
            color=texte,
        )

    feuille["B6"].number_format = "dd/mm/yyyy hh:mm:ss"
    feuille["B7"].number_format = "dd/mm/yyyy hh:mm:ss"

    ligne_entete = 10
    bordure_entete = Side(
        style="medium",
        color=orange,
    )

    for colonne, entete in enumerate(
        ["Matricule", "Nom & Prenom", "Date", "Heure"],
        start=1,
    ):
        cellule = feuille.cell(
            row=ligne_entete,
            column=colonne,
            value=entete,
        )
        cellule.font = Font(
            name="Aptos",
            bold=True,
            color=blanc,
        )
        cellule.fill = PatternFill(
            fill_type="solid",
            fgColor=bleu,
        )
        cellule.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
        cellule.border = Border(
            bottom=bordure_entete,
        )

    bordure_ligne = Side(
        style="thin",
        color=ligne,
    )

    for numero_ligne, pointage in enumerate(
        pointages,
        start=ligne_entete + 1,
    ):
        matricule = str(
            pointage.get("matricule", "") or ""
        ).strip()

        if matricule.isdigit():
            cellule_matricule = feuille.cell(
                row=numero_ligne,
                column=1,
                value=int(matricule),
            )
            cellule_matricule.number_format = "0"
        else:
            feuille.cell(
                row=numero_ligne,
                column=1,
                value=matricule,
            )

        feuille.cell(
            row=numero_ligne,
            column=2,
            value=str(pointage.get("nom", "") or "").strip(),
        )

        timestamp = pointage.get("timestamp")
        if isinstance(timestamp, datetime):
            cellule_date = feuille.cell(
                row=numero_ligne,
                column=3,
                value=timestamp.date(),
            )
            cellule_date.number_format = "dd/mm/yyyy"

            cellule_heure = feuille.cell(
                row=numero_ligne,
                column=4,
                value=timestamp.time(),
            )
            cellule_heure.number_format = "hh:mm:ss"

        couleur = blanc if numero_ligne % 2 == 0 else alternee
        remplissage = PatternFill(
            fill_type="solid",
            fgColor=couleur,
        )

        for colonne in range(1, 5):
            cellule = feuille.cell(
                row=numero_ligne,
                column=colonne,
            )
            cellule.fill = remplissage
            cellule.font = Font(
                name="Aptos",
                size=10,
                color=texte,
            )
            cellule.alignment = Alignment(
                horizontal="left" if colonne in (1, 2) else "center",
                vertical="center",
            )
            cellule.border = Border(
                bottom=bordure_ligne,
            )

    derniere_ligne = max(
        ligne_entete,
        ligne_entete + len(pointages),
    )

    feuille.auto_filter.ref = (
        f"A{ligne_entete}:D{derniere_ligne}"
    )

    # Toute la feuille peut défiler.
    feuille.freeze_panes = None

    feuille.column_dimensions["A"].width = 18
    feuille.column_dimensions["B"].width = 36
    feuille.column_dimensions["C"].width = 16
    feuille.column_dimensions["D"].width = 15

    feuille.page_setup.orientation = "landscape"
    feuille.page_setup.fitToWidth = 1
    feuille.page_setup.fitToHeight = 0
    feuille.sheet_properties.pageSetUpPr.fitToPage = True

    flux = BytesIO()

    try:
        classeur.save(flux)
    except Exception as erreur:
        classeur.close()
        return {
            "success": False,
            "contenu": None,
            "nom_fichier": None,
            "nombre": 0,
            "nombre_terminal": nombre_terminal,
            "erreur": "Impossible de creer le fichier Excel : " + str(erreur),
        }
    finally:
        try:
            classeur.close()
        except Exception:
            pass

    flux.seek(0)
    temps_excel = perf_counter() - mesure

    nom_fichier = (
        f"pointages_{_nom_fichier(nom_pointeuse)}_"
        f"{debut:%Y%m%d_%H%M}_{fin:%Y%m%d_%H%M}.xlsx"
    )

    try:
        nombre_corrections = int(
            corrections.get("nombre_corrections", 0) or 0
        )
    except (TypeError, ValueError):
        nombre_corrections = 0

    return {
        "success": True,
        "contenu": flux.getvalue(),
        "nom_fichier": nom_fichier,
        "nombre": len(pointages),
        "nombre_terminal": nombre_terminal,
        "nombre_corrections": nombre_corrections,
        "source": "sqlite",
        "performance": {
            "sqlite_secondes": round(temps_sqlite, 3),
            "corrections_secondes": round(temps_corrections, 3),
            "excel_secondes": round(temps_excel, 3),
            "total_secondes": round(
                perf_counter() - debut_total,
                3,
            ),
        },
        "erreur": None,
    }
