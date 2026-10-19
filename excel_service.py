from pathlib import Path

from openpyxl import load_workbook


# ============================================================
# CONFIGURATION
# ============================================================

LONGUEUR_MAX_MATRICULE = 5
LONGUEUR_MIN_BADGE = 6


# ============================================================
# NORMALISER LE NOM D'UNE COLONNE
# ============================================================

def normaliser_colonne(valeur):
    if valeur is None:
        return ""

    texte = str(valeur).strip().lower()

    remplacements = {
        "é": "e",
        "è": "e",
        "ê": "e",
        "ë": "e",
        "à": "a",
        "â": "a",
        "ä": "a",
        "î": "i",
        "ï": "i",
        "ô": "o",
        "ö": "o",
        "ù": "u",
        "û": "u",
        "ü": "u",
        "ç": "c",
        "’": "'",
    }

    for ancien, nouveau in remplacements.items():
        texte = texte.replace(
            ancien,
            nouveau,
        )

    texte = texte.replace(
        "°",
        "",
    )

    texte = texte.replace(
        "º",
        "",
    )

    texte = " ".join(
        texte.split()
    )

    return texte


# ============================================================
# NORMALISER UNE VALEUR EXCEL
# ============================================================

def normaliser_valeur(valeur):
    if valeur is None:
        return ""

    # OpenPyXL peut retourner :
    #
    # 5089.0
    #
    # alors que la valeur réelle est :
    #
    # 5089

    if isinstance(
        valeur,
        float,
    ):
        if valeur.is_integer():
            return str(
                int(valeur)
            )

    return str(
        valeur
    ).strip()


# ============================================================
# NORMALISER UNE VALEUR NUMERIQUE EXCEL
# ============================================================

def normaliser_nombre_excel(valeur):
    texte = normaliser_valeur(
        valeur
    )

    if not texte:
        return ""

    # Les espaces ne doivent pas empêcher
    # la reconnaissance d'une valeur numérique.

    texte_sans_espaces = (
        texte.replace(
            " ",
            "",
        )
    )

    # Exemple :
    #
    # "5089.0" -> "5089"

    if texte_sans_espaces.endswith(
        ".0"
    ):
        candidat = (
            texte_sans_espaces[:-2]
        )

        if candidat.isdigit():
            return candidat

    return texte_sans_espaces


# ============================================================
# TROUVER UNE COLONNE
# ============================================================

def trouver_colonne(
    colonnes,
    variantes,
):
    for variante in variantes:
        variante_normalisee = (
            normaliser_colonne(
                variante
            )
        )

        if variante_normalisee in colonnes:
            return colonnes[
                variante_normalisee
            ]

    return None


# ============================================================
# VARIANTES MATRICULE
# ============================================================

VARIANTES_MATRICULE = [
    "matricule",
    "matricule employe",
    "matricule employé",
    "matricule salarié",
    "matricule salarie",
    "id employe",
    "id employé",
]


# ============================================================
# VARIANTES NOM / PRENOM
# ============================================================

VARIANTES_NOM = [
    "nom",

    "nom prenom",
    "nom prénom",

    "nom & prenom",
    "nom & prénom",

    "nom et prenom",
    "nom et prénom",

    "nom / prenom",
    "nom / prénom",

    "nom-prenom",
    "nom-prénom",

    "nom complet",

    "nom et prenoms",
    "nom et prénoms",

    "nom & prenoms",
    "nom & prénoms",
]


# ============================================================
# VARIANTES BADGE / CARTE
# ============================================================

VARIANTES_BADGE = [
    "numero de badge",
    "numéro de badge",

    "numero badge",
    "numéro badge",

    "num badge",

    "n badge",
    "n° badge",

    "badge",

    "badge numero",
    "badge numéro",

    "numero de carte",
    "numéro de carte",

    "numero carte",
    "numéro carte",

    "num carte",

    "n carte",
    "n° carte",

    "carte",

    "carte numero",
    "carte numéro",

    "numero de badge ou carte",
    "numéro de badge ou carte",

    "numero badge ou carte",
    "numéro badge ou carte",

    "badge ou carte",

    "badge / carte",
    "badge/carte",

    "numero de badge / carte",
    "numéro de badge / carte",

    "numero badge / carte",
    "numéro badge / carte",

    "numero de carte ou badge",
    "numéro de carte ou badge",

    "carte ou badge",
]


# ============================================================
# DETECTER LES COLONNES CONNUES
# ============================================================

def detecter_colonnes_entete(
    premiere_ligne,
):
    colonnes = {}

    for index, valeur in enumerate(
        premiere_ligne
    ):
        nom_colonne = (
            normaliser_colonne(
                valeur
            )
        )

        if nom_colonne:
            colonnes[
                nom_colonne
            ] = index

    colonne_matricule = trouver_colonne(
        colonnes,
        VARIANTES_MATRICULE,
    )

    colonne_nom = trouver_colonne(
        colonnes,
        VARIANTES_NOM,
    )

    colonne_badge = trouver_colonne(
        colonnes,
        VARIANTES_BADGE,
    )

    complet = (
        colonne_matricule is not None
        and colonne_nom is not None
        and colonne_badge is not None
    )

    partiel = any(
        colonne is not None
        for colonne in (
            colonne_matricule,
            colonne_nom,
            colonne_badge,
        )
    )

    return {
        "complet": complet,
        "partiel": partiel,
        "matricule": colonne_matricule,
        "nom": colonne_nom,
        "badge": colonne_badge,
    }


# ============================================================
# CLASSER AUTOMATIQUEMENT UNE CELLULE
#
# REGLES :
#
# 1 a 5 chiffres
#       -> MATRICULE
#
# 6 chiffres ou plus
#       -> BADGE / CARTE
#
# Texte
#       -> NOM / PRENOM
# ============================================================

def classifier_valeur_excel(
    valeur,
):
    texte = normaliser_nombre_excel(
        valeur
    )

    if not texte:
        return None, ""

    # ========================================================
    # NUMERIQUE
    # ========================================================

    if texte.isdigit():
        longueur = len(
            texte
        )

        # ----------------------------------------------------
        # MATRICULE
        # ----------------------------------------------------

        if (
            longueur
            <= LONGUEUR_MAX_MATRICULE
        ):
            return (
                "matricule",
                texte,
            )

        # ----------------------------------------------------
        # BADGE / CARTE
        # ----------------------------------------------------

        if (
            longueur
            >= LONGUEUR_MIN_BADGE
        ):
            return (
                "badge",
                texte,
            )

    # ========================================================
    # TEXTE
    # ========================================================

    return (
        "texte",
        normaliser_valeur(
            valeur
        ),
    )


# ============================================================
# EXTRAIRE UNE LIGNE AVEC EN-TETES
# ============================================================

def extraire_ligne_guidee(
    ligne,
    colonnes,
):
    def valeur_index(index):
        if (
            index is None
            or index >= len(ligne)
        ):
            return ""

        return normaliser_valeur(
            ligne[index]
        )

    return {
        "matricule": valeur_index(
            colonnes["matricule"]
        ),

        "nom": valeur_index(
            colonnes["nom"]
        ),

        "badge": valeur_index(
            colonnes["badge"]
        ),

        "ambiguities": [],
    }


# ============================================================
# EXTRAIRE UNE LIGNE AUTOMATIQUEMENT
# ============================================================

def extraire_ligne_automatique(
    ligne,
):
    matricules = []
    badges = []
    textes = []

    for valeur in ligne:
        (
            type_valeur,
            contenu,
        ) = classifier_valeur_excel(
            valeur
        )

        if not contenu:
            continue

        # ----------------------------------------------------
        # MATRICULE
        # ----------------------------------------------------

        if type_valeur == "matricule":
            matricules.append(
                contenu
            )

        # ----------------------------------------------------
        # BADGE
        # ----------------------------------------------------

        elif type_valeur == "badge":
            badges.append(
                contenu
            )

        # ----------------------------------------------------
        # TEXTE
        # ----------------------------------------------------

        elif type_valeur == "texte":
            textes.append(
                contenu
            )

    ambiguities = []

    # ========================================================
    # PLUSIEURS MATRICULES POSSIBLES
    # ========================================================

    if len(matricules) > 1:
        ambiguities.append(
            (
                "Plusieurs matricules possibles : "
                + ", ".join(
                    matricules
                )
                + "."
            )
        )

    # ========================================================
    # PLUSIEURS BADGES POSSIBLES
    # ========================================================

    if len(badges) > 1:
        ambiguities.append(
            (
                "Plusieurs badges / cartes possibles : "
                + ", ".join(
                    badges
                )
                + "."
            )
        )

    # ========================================================
    # NOM + PRENOM
    #
    # Si le fichier contient :
    #
    # NOM | PRENOM
    #
    # Exemple :
    #
    # TOUKEBRI | SAIFEDDINE
    #
    # les deux cellules texte sont concatenees :
    #
    # TOUKEBRI SAIFEDDINE
    # ========================================================

    nom = " ".join(
        texte
        for texte in textes
        if texte
    ).strip()

    return {
        "matricule": (
            matricules[0]
            if len(matricules) == 1
            else ""
        ),

        "nom": nom,

        "badge": (
            badges[0]
            if len(badges) == 1
            else ""
        ),

        "ambiguities": ambiguities,
    }


# ============================================================
# VALIDER UN EMPLOYE DETECTE
# ============================================================

def valider_employe_detecte(
    employe,
):
    matricule = (
        normaliser_nombre_excel(
            employe.get(
                "matricule"
            )
        )
    )

    nom = normaliser_valeur(
        employe.get(
            "nom"
        )
    )

    badge = normaliser_nombre_excel(
        employe.get(
            "badge"
        )
    )

    erreurs = list(
        employe.get(
            "ambiguities",
            [],
        )
    )

    # ========================================================
    # MATRICULE
    # ========================================================

    if not matricule:
        erreurs.append(
            (
                "Matricule manquant "
                "ou non identifiable."
            )
        )

    elif not matricule.isdigit():
        erreurs.append(
            (
                "Le matricule doit contenir "
                "uniquement des chiffres."
            )
        )

    elif (
        len(matricule)
        > LONGUEUR_MAX_MATRICULE
    ):
        erreurs.append(
            (
                "Le matricule doit contenir "
                "au maximum "
                + str(
                    LONGUEUR_MAX_MATRICULE
                )
                + " chiffres."
            )
        )

    # ========================================================
    # NOM
    # ========================================================

    if not nom:
        erreurs.append(
            (
                "Nom & Prénom manquant "
                "ou non identifiable."
            )
        )

    # ========================================================
    # BADGE
    # ========================================================

    if not badge:
        erreurs.append(
            (
                "Numéro de badge / carte "
                "manquant ou non identifiable."
            )
        )

    elif not badge.isdigit():
        erreurs.append(
            (
                "Le numéro de badge / carte "
                "doit contenir uniquement "
                "des chiffres."
            )
        )

    elif (
        len(badge)
        < LONGUEUR_MIN_BADGE
    ):
        erreurs.append(
            (
                "Le numéro de badge / carte "
                "doit contenir au moins "
                + str(
                    LONGUEUR_MIN_BADGE
                )
                + " chiffres."
            )
        )

    return {
        "matricule": matricule,
        "nom": nom,
        "badge": badge,
        "erreurs": erreurs,
    }


# ============================================================
# LIRE UN FICHIER EXCEL D'EMPLOYES
# ============================================================

def lire_fichier_employes_excel(
    chemin_fichier,
):
    chemin = Path(
        chemin_fichier
    )

    # ========================================================
    # VERIFIER EXISTENCE
    # ========================================================

    if not chemin.exists():
        return {
            "success": False,
            "employes": [],
            "nombre": 0,

            "erreurs": [
                "Le fichier Excel est introuvable."
            ],

            "nombre_erreurs": 1,

            "mode_detection": None,
        }

    # ========================================================
    # VERIFIER EXTENSION
    # ========================================================

    if chemin.suffix.lower() != ".xlsx":
        return {
            "success": False,
            "employes": [],
            "nombre": 0,

            "erreurs": [
                (
                    "Le fichier doit être "
                    "au format .xlsx."
                )
            ],

            "nombre_erreurs": 1,

            "mode_detection": None,
        }

    # ========================================================
    # OUVRIR EXCEL
    # ========================================================

    try:
        classeur = load_workbook(
            filename=chemin,
            read_only=True,
            data_only=True,
        )

    except Exception as erreur:
        return {
            "success": False,
            "employes": [],
            "nombre": 0,

            "erreurs": [
                (
                    "Impossible de lire "
                    "le fichier Excel : "
                    + str(erreur)
                )
            ],

            "nombre_erreurs": 1,

            "mode_detection": None,
        }

    try:
        feuille = classeur.active

        # ====================================================
        # PREMIERE LIGNE
        # ====================================================

        premiere_ligne = next(
            feuille.iter_rows(
                min_row=1,
                max_row=1,
                values_only=True,
            ),
            None,
        )

        if premiere_ligne is None:
            return {
                "success": False,
                "employes": [],
                "nombre": 0,

                "erreurs": [
                    "Le fichier Excel est vide."
                ],

                "nombre_erreurs": 1,

                "mode_detection": None,
            }

        # ====================================================
        # DETECTION DES EN-TETES
        # ====================================================

        detection_entete = (
            detecter_colonnes_entete(
                premiere_ligne
            )
        )

        # ====================================================
        # MODE 1
        #
        # LES 3 EN-TETES SONT RECONNUS
        #
        # La lecture guidée reste prioritaire.
        # ====================================================

        if detection_entete[
            "complet"
        ]:
            mode_detection = (
                "entetes"
            )

            ligne_depart = 2

        # ====================================================
        # MODE 2
        #
        # AUCUN EN-TETE RECONNU
        #
        # La première ligne est considérée
        # comme une vraie ligne employé.
        # ====================================================

        elif not detection_entete[
            "partiel"
        ]:
            mode_detection = (
                "automatique"
            )

            ligne_depart = 1

        # ====================================================
        # MODE 3
        #
        # EN-TETE PARTIEL
        #
        # Exemple :
        #
        # Matricule | Colonne B | Colonne C
        #
        # La première ligne est considérée comme
        # un en-tête, mais les lignes suivantes
        # sont analysées automatiquement.
        # ====================================================

        else:
            mode_detection = (
                "automatique_apres_entete_partiel"
            )

            ligne_depart = 2

        # ====================================================
        # RESULTATS
        # ====================================================

        employes = []
        erreurs = []

        matricules_excel = set()
        badges_excel = set()

        # ====================================================
        # PARCOURIR LES LIGNES
        # ====================================================

        for numero_ligne, ligne in enumerate(
            feuille.iter_rows(
                min_row=ligne_depart,
                values_only=True,
            ),
            start=ligne_depart,
        ):
            if not ligne:
                continue

            valeurs_ligne = [
                normaliser_valeur(
                    valeur
                )
                for valeur in ligne
            ]

            # ------------------------------------------------
            # IGNORER LIGNE VIDE
            # ------------------------------------------------

            if not any(
                valeurs_ligne
            ):
                continue

            # =================================================
            # EXTRACTION
            # =================================================

            if (
                mode_detection
                == "entetes"
            ):
                extrait = (
                    extraire_ligne_guidee(
                        ligne,
                        detection_entete,
                    )
                )

            else:
                extrait = (
                    extraire_ligne_automatique(
                        ligne
                    )
                )

            # =================================================
            # VALIDATION
            # =================================================

            valide = (
                valider_employe_detecte(
                    extrait
                )
            )

            matricule = valide[
                "matricule"
            ]

            nom = valide[
                "nom"
            ]

            badge = valide[
                "badge"
            ]

            erreurs_ligne = list(
                valide[
                    "erreurs"
                ]
            )

            # =================================================
            # DOUBLON MATRICULE
            # =================================================

            cle_matricule = (
                matricule.upper()
                if matricule
                else ""
            )

            if matricule:
                if (
                    cle_matricule
                    in matricules_excel
                ):
                    erreurs_ligne.append(
                        (
                            "Matricule dupliqué "
                            "dans le fichier Excel."
                        )
                    )

                else:
                    matricules_excel.add(
                        cle_matricule
                    )

            # =================================================
            # DOUBLON BADGE
            # =================================================

            if badge:
                if (
                    badge
                    in badges_excel
                ):
                    erreurs_ligne.append(
                        (
                            "Badge / numéro de carte "
                            "dupliqué dans le fichier Excel."
                        )
                    )

                else:
                    badges_excel.add(
                        badge
                    )

            # =================================================
            # LIGNE INVALIDE
            # =================================================

            if erreurs_ligne:
                erreurs.append(
                    {
                        "ligne": numero_ligne,
                        "matricule": matricule,
                        "nom": nom,
                        "badge": badge,
                        "erreurs": erreurs_ligne,
                    }
                )

                continue

            # =================================================
            # LIGNE VALIDE
            # =================================================

            employes.append(
                {
                    "ligne": numero_ligne,
                    "matricule": matricule,
                    "nom": nom,
                    "badge": badge,
                }
            )

        # ====================================================
        # RESULTAT FINAL
        # ====================================================

        return {
            "success": (
                len(erreurs) == 0
            ),

            "employes": employes,

            "nombre": len(
                employes
            ),

            "erreurs": erreurs,

            "nombre_erreurs": len(
                erreurs
            ),

            "mode_detection": (
                mode_detection
            ),

            "regles_detection": {
                "matricule": (
                    "1 à 5 chiffres"
                ),

                "nom": (
                    "texte"
                ),

                "badge": (
                    "6 chiffres ou plus"
                ),
            },
        }

    except Exception as erreur:
        return {
            "success": False,

            "employes": [],

            "nombre": 0,

            "erreurs": [
                (
                    "Erreur pendant l'analyse "
                    "du fichier Excel : "
                    + str(erreur)
                )
            ],

            "nombre_erreurs": 1,

            "mode_detection": None,
        }

    finally:
        classeur.close()